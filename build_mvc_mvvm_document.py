from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("artifacts/mvc_mvvm")
OUT_DIR.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT_DIR / "MVC_und_MVVM_Datenfluss.docx"
MVC_PNG = OUT_DIR / "mvc_datenfluss.png"
MVVM_PNG = OUT_DIR / "mvvm_datenfluss.png"

NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "0F6B78"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F5F5"
INK = "1F2933"
MUTED = "52616B"
LINE = "D9E2F3"
WHITE = "FFFFFF"


def paint(value):
    """Return a Pillow-compatible color while keeping OOXML colors hash-free."""
    return value if value.startswith("#") else "#" + value


def font(size, bold=False):
    paths = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in paths:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def rounded_box(draw, xy, title, subtitle, fill, title_font, text_font):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill=paint(fill), outline=paint(NAVY), width=3)
    title_box = draw.textbbox((0, 0), title, font=title_font)
    sub_box = draw.textbbox((0, 0), subtitle, font=text_font)
    draw.text(((x1 + x2 - (title_box[2] - title_box[0])) / 2, y1 + 27), title, font=title_font, fill=paint(INK))
    draw.text(((x1 + x2 - (sub_box[2] - sub_box[0])) / 2, y1 + 82), subtitle, font=text_font, fill=paint(MUTED))


def arrow(draw, start, end, label, label_pos, arrow_color=BLUE, label_color=MUTED, width=5):
    draw.line([start, end], fill=paint(arrow_color), width=width)
    sx, sy = start
    ex, ey = end
    # The arrows in these diagrams are horizontal or vertical. Draw a short, clear head.
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        head = [(ex, ey), (ex - 18 * direction, ey - 11), (ex - 18 * direction, ey + 11)]
    else:
        direction = 1 if ey > sy else -1
        head = [(ex, ey), (ex - 11, ey - 18 * direction), (ex + 11, ey - 18 * direction)]
    draw.polygon(head, fill=paint(arrow_color))
    label_font = font(24)
    bbox = draw.textbbox((0, 0), label, font=label_font)
    lx, ly = label_pos
    pad_x, pad_y = 10, 5
    draw.rounded_rectangle(
        (lx - pad_x, ly - pad_y, lx + bbox[2] - bbox[0] + pad_x, ly + bbox[3] - bbox[1] + pad_y),
        radius=8, fill=paint(WHITE)
    )
    draw.text((lx, ly), label, font=label_font, fill=paint(label_color))


def connector(draw, points, label, label_pos, color=BLUE):
    for a, b in zip(points, points[1:]):
        draw.line([a, b], fill=paint(color), width=5)
    # Arrow head on the final segment.
    arrow(draw, points[-2], points[-1], label, label_pos, arrow_color=color, width=5)


def make_mvc_diagram(path):
    image = Image.new("RGB", (1600, 610), paint(WHITE))
    draw = ImageDraw.Draw(image)
    title_font = font(34, bold=True)
    sub_font = font(23)
    node_title = font(30, bold=True)
    node_text = font(21)
    draw.text((55, 26), "MVC: Der Controller verarbeitet die Eingabe", font=title_font, fill=paint(INK))
    draw.text((55, 77), "Typischer Ablauf: Eingabe wird verarbeitet, das Modell ändert sich, die Ansicht wird neu dargestellt.", font=sub_font, fill=paint(MUTED))

    rounded_box(draw, (70, 170, 340, 315), "Nutzende", "klicken oder tippen", "F8FBFE", node_title, node_text)
    rounded_box(draw, (445, 170, 715, 315), "View", "Oberfläche", PALE_BLUE, node_title, node_text)
    rounded_box(draw, (820, 170, 1090, 315), "Controller", "wertet Eingabe aus", "E8F0F8", node_title, node_text)
    rounded_box(draw, (1195, 170, 1465, 315), "Model", "Daten und Logik", "E9F6F7", node_title, node_text)

    arrow(draw, (340, 242), (445, 242), "1 Eingabe", (348, 197))
    arrow(draw, (715, 242), (820, 242), "2 Event", (727, 197))
    arrow(draw, (1090, 242), (1195, 242), "3 ändern oder lesen", (1040, 197))
    connector(draw, [(1330, 315), (1330, 450), (580, 450), (580, 315)], "4 neuer Zustand", (865, 414), color=TEAL)
    connector(draw, [(445, 492), (205, 492), (205, 315)], "5 aktualisierte Ansicht", (220, 510), color=TEAL)
    image.save(path)


def make_mvvm_diagram(path):
    image = Image.new("RGB", (1600, 650), paint(WHITE))
    draw = ImageDraw.Draw(image)
    title_font = font(34, bold=True)
    sub_font = font(23)
    node_title = font(30, bold=True)
    node_text = font(21)
    draw.text((55, 26), "MVVM: Die View ist über Data Binding mit dem ViewModel verbunden", font=title_font, fill=paint(INK))
    draw.text((55, 77), "Typischer Ablauf: Die Oberfläche sendet Befehle; sichtbare Werte werden automatisch abgeglichen.", font=sub_font, fill=paint(MUTED))

    rounded_box(draw, (70, 185, 340, 330), "Nutzende", "klicken oder tippen", "F8FBFE", node_title, node_text)
    rounded_box(draw, (445, 185, 715, 330), "View", "Oberfläche", PALE_BLUE, node_title, node_text)
    rounded_box(draw, (820, 185, 1090, 330), "ViewModel", "Zustand und Befehle", "E8F0F8", node_title, node_text)
    rounded_box(draw, (1195, 185, 1465, 330), "Model", "Daten und Logik", PALE_TEAL, node_title, node_text)

    arrow(draw, (340, 257), (445, 257), "1 Eingabe", (350, 210))
    arrow(draw, (715, 257), (820, 257), "2 Befehl", (727, 210))
    arrow(draw, (1090, 257), (1195, 257), "3 lesen oder ändern", (1040, 210))
    connector(draw, [(1330, 330), (1330, 470), (955, 470), (955, 330)], "4 Daten und Zustand", (1085, 434), color=TEAL)
    # Bidirectional data binding is represented by two opposite, parallel arrows.
    arrow(draw, (820, 385), (715, 385), "5 Werte zur View", (560, 347), arrow_color=TEAL)
    arrow(draw, (715, 440), (820, 440), "Data Binding und Änderungen", (500, 455), arrow_color=TEAL)
    connector(draw, [(445, 520), (205, 520), (205, 330)], "6 aktualisierte Oberfläche", (220, 540), color=TEAL)
    image.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_run(paragraph, text, size=11, bold=False, color=INK, italic=False):
    run = paragraph.add_run(text)
    run.font.name = "Liberation Sans"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def style_paragraph(paragraph, before=0, after=0, line=1.15, alignment=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment


def add_figure(doc, image_path, caption, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=7, after=2)
    p.add_run().add_picture(str(image_path), width=Inches(6.85))
    inline_shape = doc.inline_shapes[-1]
    inline_shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(cap, before=0, after=9, line=1.0)
    add_run(cap, caption, size=9, italic=True, color=MUTED)


def add_bullet(doc, lead, rest):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=3, line=1.08)
    add_run(p, lead, size=10.5, bold=True)
    add_run(p, rest, size=10.5)
    return p


def build_document():
    make_mvc_diagram(MVC_PNG)
    make_mvvm_diagram(MVVM_PNG)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)

    title_style = doc.styles["Title"]
    title_style.font.name = "Liberation Sans"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_border(title_style)

    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Liberation Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    # Page 1
    title = doc.add_paragraph(style="Title")
    remove_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, before=0, after=2, line=1.0)
    add_run(title, "MVC und MVVM", size=25, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(subtitle, after=16, line=1.1)
    add_run(subtitle, "Unterschiede und Datenfluss in zwei Architekturmustern", size=12, color="000000")

    p = doc.add_paragraph()
    style_paragraph(p, after=10, line=1.18)
    add_run(p, "Kurz gesagt: ", size=11, bold=True)
    add_run(p, "Beide Muster trennen Daten und Benutzeroberfläche. Bei MVC verarbeitet ein Controller die Eingabe. Bei MVVM verbindet Data Binding die View mit dem ViewModel und hält sichtbare Werte synchron.", size=11)

    h1 = doc.add_paragraph(style="Heading 1")
    style_paragraph(h1, before=4, after=2)
    set_keep_with_next(h1)
    add_run(h1, "1 MVC Model View Controller", size=15, bold=True, color="000000")

    p = doc.add_paragraph()
    style_paragraph(p, after=5, line=1.15)
    add_run(p, "Model", size=11, bold=True)
    add_run(p, " enthält Daten und Fachlogik. Die ", size=11)
    add_run(p, "View", size=11, bold=True)
    add_run(p, " zeigt diese Daten an. Der ", size=11)
    add_run(p, "Controller", size=11, bold=True)
    add_run(p, " nimmt Eingaben aus der View entgegen und veranlasst Änderungen am Model. Anschließend wird die View mit dem neuen Zustand dargestellt.", size=11)

    add_figure(
        doc,
        MVC_PNG,
        "Abbildung 1: Vereinfachter MVC Datenfluss. Die genaue Umsetzung kann je nach Framework variieren.",
        "MVC Datenfluss: Nutzende senden Eingaben an die View, die den Controller informiert. Der Controller ändert oder liest das Model. Das Model liefert den neuen Zustand an die View, die die aktualisierte Ansicht zeigt.",
    )

    h2 = doc.add_paragraph(style="Heading 2")
    style_paragraph(h2, before=0, after=2)
    set_keep_with_next(h2)
    add_run(h2, "Typische Rollen", size=12, bold=True, color="000000")
    add_bullet(doc, "Model: ", "Speichert Daten und enthält Regeln der Anwendung.")
    add_bullet(doc, "View: ", "Zeigt Inhalte an und nimmt Eingaben entgegen.")
    add_bullet(doc, "Controller: ", "Entscheidet, was nach einer Eingabe passieren soll.")

    doc.add_page_break()

    # Page 2
    h1 = doc.add_paragraph(style="Heading 1")
    style_paragraph(h1, before=0, after=2)
    set_keep_with_next(h1)
    add_run(h1, "2 MVVM Model View ViewModel", size=15, bold=True, color="000000")

    p = doc.add_paragraph()
    style_paragraph(p, after=5, line=1.15)
    add_run(p, "Das Model bleibt für Daten und Fachlogik zuständig. Die ", size=11)
    add_run(p, "ViewModel", size=11, bold=True)
    add_run(p, " stellt der View die Werte und Befehle bereit, die sie braucht. Die View kennt die Darstellung, aber nicht die Geschäftslogik. Über ", size=11)
    add_run(p, "Data Binding", size=11, bold=True)
    add_run(p, " werden Änderungen zwischen View und ViewModel automatisch abgeglichen.", size=11)

    add_figure(
        doc,
        MVVM_PNG,
        "Abbildung 2: Vereinfachter MVVM Datenfluss. Data Binding ersetzt nicht das Model, sondern verbindet View und ViewModel.",
        "MVVM Datenfluss: Nutzende bedienen die View. Die View übergibt Befehle an das ViewModel, das mit dem Model arbeitet. Data Binding hält Werte zwischen View und ViewModel synchron und die View zeigt die aktualisierte Oberfläche.",
    )

    h1 = doc.add_paragraph(style="Heading 1")
    style_paragraph(h1, before=2, after=4)
    set_keep_with_next(h1)
    add_run(h1, "3 Der wichtigste Unterschied", size=15, bold=True, color="000000")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Frage", "MVC", "MVVM"]
    widths = [1.35, 2.65, 2.65]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_col_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, before=0, after=0, line=1.05, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_run(p, text, size=10, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])

    rows = [
        ("Eingabe", "View gibt Ereignis an den Controller weiter.", "View ruft einen Befehl im ViewModel auf."),
        ("Aktualisierung", "Der Controller und die View reagieren auf Änderungen am Model.", "Data Binding gleicht Werte zwischen View und ViewModel ab."),
        ("Kernidee", "Ein Controller vermittelt zwischen Eingabe, Model und Anzeige.", "Ein ViewModel beschreibt den Zustand der Oberfläche ohne direkte View Referenz."),
        ("Häufig passend für", "Webanwendungen mit klaren Anfrage Antwort Abläufen.", "Reaktive Oberflächen mit vielen automatisch aktualisierten Feldern."),
    ]
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "F6F9FC" if r_idx % 2 == 0 else WHITE
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_col_width(cell, widths[idx])
            set_cell_shading(cell, fill)
            set_cell_border(cell)
            set_cell_margin(cell, top=105, start=120, bottom=105, end=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, before=0, after=0, line=1.08, alignment=WD_ALIGN_PARAGRAPH.LEFT)
            add_run(p, text, size=9.5, bold=(idx == 0), color=INK)

    p = doc.add_paragraph()
    style_paragraph(p, before=10, after=0, line=1.15)
    add_run(p, "Merksatz: ", size=11, bold=True)
    add_run(p, "MVC betont die Steuerung durch den Controller. MVVM betont die Bindung zwischen Oberfläche und ViewModel.", size=11)

    doc.core_properties.title = "MVC und MVVM"
    doc.core_properties.subject = "Unterschiede und Datenfluss in MVC und MVVM"
    doc.core_properties.author = ""
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
    print(DOCX_PATH.resolve())
